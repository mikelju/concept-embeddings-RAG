#!/usr/bin/env python3
"""Convert one or more Markdown files to Word (.docx) format.

Usage:
    python md_to_docx.py archivo1.md archivo2.md
    python md_to_docx.py archivo.md --output-dir carpeta_salida/

Dependencies: python-docx, markdown-it-py
"""

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from markdown_it import MarkdownIt


# ---------------------------------------------------------------------------
# Inline content rendering
# ---------------------------------------------------------------------------

def _add_inline_runs(paragraph, inline_token):
    """Walk inline token children and add formatted runs to a paragraph."""
    if not inline_token.children:
        return

    bold = False
    italic = False

    for child in inline_token.children:
        if child.type == "strong_open":
            bold = True
        elif child.type == "strong_close":
            bold = False
        elif child.type == "em_open":
            italic = True
        elif child.type == "em_close":
            italic = False
        elif child.type == "text":
            run = paragraph.add_run(child.content)
            run.bold = bold
            run.italic = italic
        elif child.type == "code_inline":
            run = paragraph.add_run(child.content)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x58, 0x5E, 0x68)
        elif child.type in ("softbreak", "hardbreak"):
            paragraph.add_run("\n")
        # link_open / link_close / image — render text only


# ---------------------------------------------------------------------------
# Main converter
# ---------------------------------------------------------------------------

def convert_md_to_docx(md_path: Path, output_dir: Path | None = None) -> Path:
    """Convert a single Markdown file to .docx and return the output path."""

    md_text = md_path.read_text(encoding="utf-8")
    md = MarkdownIt("commonmark", {"typographer": True}).enable("table")
    tokens = md.parse(md_text)

    doc = Document()

    # Base style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # State
    list_stack: list[str] = []     # "bullet" | "ordered"
    in_blockquote = False
    table_rows: list[list[tuple]] | None = None
    current_row: list[tuple] | None = None
    header_row = False

    i = 0
    while i < len(tokens):
        tok = tokens[i]

        # -- Headings
        if tok.type == "heading_open":
            level = int(tok.tag[1])
            i += 1
            heading = doc.add_heading(level=min(level, 9))
            if i < len(tokens) and tokens[i].type == "inline":
                _add_inline_runs(heading, tokens[i])
                i += 1
            i += 1  # heading_close
            continue

        # -- Paragraphs
        if tok.type == "paragraph_open":
            i += 1
            inline_tok = None
            if i < len(tokens) and tokens[i].type == "inline":
                inline_tok = tokens[i]
                i += 1
            i += 1  # paragraph_close

            if list_stack:
                style_name = "List Bullet" if list_stack[-1] == "bullet" else "List Number"
                p = doc.add_paragraph(style=style_name)
                if len(list_stack) > 1:
                    p.paragraph_format.left_indent = Cm(1.27 * (len(list_stack) - 1))
            elif in_blockquote:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.0)
            else:
                p = doc.add_paragraph()

            if inline_tok:
                _add_inline_runs(p, inline_tok)
                if in_blockquote:
                    for run in p.runs:
                        run.italic = True
                        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            continue

        # -- Lists
        if tok.type == "bullet_list_open":
            list_stack.append("bullet")
            i += 1
            continue
        if tok.type == "bullet_list_close":
            list_stack.pop()
            i += 1
            continue
        if tok.type == "ordered_list_open":
            list_stack.append("ordered")
            i += 1
            continue
        if tok.type == "ordered_list_close":
            list_stack.pop()
            i += 1
            continue
        if tok.type in ("list_item_open", "list_item_close"):
            i += 1
            continue

        # -- Blockquotes
        if tok.type == "blockquote_open":
            in_blockquote = True
            i += 1
            continue
        if tok.type == "blockquote_close":
            in_blockquote = False
            i += 1
            continue

        # -- Code blocks
        if tok.type in ("fence", "code_block"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            run = p.add_run(tok.content.rstrip())
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
            i += 1
            continue

        # -- Horizontal rule
        if tok.type == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("\u2500" * 60)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            run.font.size = Pt(8)
            i += 1
            continue

        # -- Tables
        if tok.type == "table_open":
            table_rows = []
            i += 1
            continue
        if tok.type in ("thead_open", "tbody_open", "tbody_close"):
            if tok.type == "thead_open":
                header_row = True
            i += 1
            continue
        if tok.type == "thead_close":
            header_row = False
            i += 1
            continue
        if tok.type == "tr_open":
            current_row = []
            i += 1
            continue
        if tok.type == "tr_close":
            if table_rows is not None and current_row is not None:
                table_rows.append((current_row, header_row))
            current_row = None
            i += 1
            continue
        if tok.type in ("th_open", "td_open"):
            is_header_cell = tok.type == "th_open"
            i += 1
            cell_inline = None
            if i < len(tokens) and tokens[i].type == "inline":
                cell_inline = tokens[i]
                i += 1
            if i < len(tokens) and tokens[i].type in ("th_close", "td_close"):
                i += 1
            if current_row is not None:
                current_row.append((cell_inline, is_header_cell))
            continue
        if tok.type in ("th_close", "td_close"):
            i += 1
            continue
        if tok.type == "table_close":
            if table_rows:
                num_cols = max(len(cells) for cells, _ in table_rows)
                tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
                tbl.style = "Table Grid"
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, (cells, _is_hdr) in enumerate(table_rows):
                    for c_idx, (inline_tok, is_th) in enumerate(cells):
                        if c_idx >= num_cols:
                            break
                        cell = tbl.rows[r_idx].cells[c_idx]
                        cell.text = ""
                        p = cell.paragraphs[0]
                        if inline_tok:
                            _add_inline_runs(p, inline_tok)
                        if _is_hdr or is_th:
                            for run in p.runs:
                                run.bold = True
            table_rows = None
            i += 1
            continue

        # -- Fallback
        i += 1

    # Save
    if output_dir:
        output_dir.mkdir(parents=True, exist_ok=True)
        out_path = output_dir / (md_path.stem + ".docx")
    else:
        out_path = md_path.with_suffix(".docx")

    doc.save(str(out_path))
    return out_path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Convert Markdown files to Word (.docx)"
    )
    parser.add_argument("files", nargs="+", help="Markdown file(s) to convert")
    parser.add_argument(
        "-o", "--output-dir", default=None,
        help="Output directory (default: same folder as each input file)"
    )
    args = parser.parse_args()

    ok, fail = 0, 0
    for f in args.files:
        p = Path(f)
        if not p.exists():
            print(f"  SKIP  {f} (not found)", file=sys.stderr)
            fail += 1
            continue
        try:
            out = convert_md_to_docx(p, Path(args.output_dir) if args.output_dir else None)
            print(f"  OK    {p.name}  ->  {out}")
            ok += 1
        except Exception as exc:
            print(f"  FAIL  {p.name}: {exc}", file=sys.stderr)
            fail += 1

    print(f"\nDone: {ok} converted, {fail} failed.")
    return 0 if fail == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
